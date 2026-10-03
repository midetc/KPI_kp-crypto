from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from crypto import TrithemiusCipher, UA_ALPHABET


def set_run(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_center(doc, text, size=14, bold=True, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)


def add_p(doc, text, size=12, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(6)


def add_shot_box(doc, title, what_to_do):
    add_p(doc, title, bold=True)
    add_p(doc, what_to_do, color=RGBColor(0xC0, 0x00, 0x00))
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.width = Cm(16)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:color"), "C00000")
        tcBorders.append(border)
    tcPr.append(tcBorders)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "\n\n\n\n"
        ">>> ВСТАВ СЮДИ PRINT SCREEN <<<\n"
        "(Ctrl+V після PrtSc / Win+Shift+S)\n"
        "\n\n\n"
    )
    set_run(r, size=11, bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    doc.add_paragraph()


def main():
    cipher = TrithemiusCipher()
    plain = "Алфавіт"
    linear_key = "2,5"
    nonlin_key = "1,2,3"
    motto = "ключ"
    enc_lin = cipher.encrypt(plain, "linear", linear_key, UA_ALPHABET)
    dec_lin = cipher.decrypt(enc_lin, "linear", linear_key, UA_ALPHABET)
    enc_nl = cipher.encrypt(plain, "nonlinear", nonlin_key, UA_ALPHABET)
    dec_nl = cipher.decrypt(enc_nl, "nonlinear", nonlin_key, UA_ALPHABET)
    enc_m = cipher.encrypt(plain, "motto", motto, UA_ALPHABET)
    dec_m = cipher.decrypt(enc_m, "motto", motto, UA_ALPHABET)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    add_center(doc, "НАЦІОНАЛЬНИЙ ТЕХНІЧНИЙ УНІВЕРСИТЕТ УКРАЇНИ", size=12, bold=True)
    add_center(doc, "«КИЇВСЬКИЙ ПОЛІТЕХНІЧНИЙ ІНСТИТУТ імені ІГОРЯ СІКОРСЬКОГО»", size=12)
    add_center(
        doc,
        "Дисципліна: Безпека інформаційних систем",
        size=12,
        color=RGBColor(0x00, 0x55, 0x99),
    )
    doc.add_paragraph()
    add_center(doc, "ЗВІТ", size=16, bold=True)
    add_center(doc, "до комп'ютерного практикуму №2", size=14, bold=True)
    add_center(doc, "Тема: Шифр Тритеміуса", size=14, bold=True)
    doc.add_paragraph()
    add_p(doc, "Студент: Михайлець Артем Миколайович")
    add_p(doc, "Група: ТВ-33")
    add_p(doc, "Викладач: ________________________________")
    add_p(doc, "м. Київ, 2026 р.")
    doc.add_page_break()

    add_center(doc, "1. Мета роботи", size=14)
    add_p(
        doc,
        "Модифікувати криптосистему з КП №1 для реалізації шифру Тритеміуса "
        "з ключами: лінійним (A, B), нелінійним (A, B, C) та гаслом.",
    )

    add_center(doc, "2. Короткі теоретичні відомості", size=14)
    add_p(
        doc,
        "Шифр Тритеміуса — удосконалення шифру Цезаря зі змінним кроком зміщення k. "
        "Формули ті самі y = (x + k) mod n, але k залежить від позиції p літери:",
    )
    add_p(doc, "лінійний: k = A·p + B", bold=True)
    add_p(doc, "нелінійний: k = A·p² + B·p + C", bold=True)
    add_p(doc, "гасло: k береться з повторюваного ключового рядка", bold=True)

    add_center(doc, "3. Тестові приклади", size=14)
    add_p(doc, f'Відкритий текст: "{plain}"')
    add_p(doc, f'Лінійний ключ {linear_key}: "{enc_lin}" → розшифр. "{dec_lin}"')
    add_p(doc, f'Нелінійний ключ {nonlin_key}: "{enc_nl}" → розшифр. "{dec_nl}"')
    add_p(doc, f'Гасло «{motto}»: "{enc_m}" → розшифр. "{dec_m}"')

    add_center(doc, "4. Результати тестування (скріншоти)", size=14)

    add_shot_box(
        doc,
        "Крок 1. Лінійний режим (A, B)",
        f"Що зробити: режим «Лінійний», ключ {linear_key}, текст «{plain}», "
        f"зашифруй (очікується «{enc_lin}»), потім розшифруй. Print Screen — у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 2. Нелінійний режим (A, B, C)",
        f"Що зробити: режим «Нелінійний», ключ {nonlin_key}, текст «{plain}», "
        f"зашифруй (очікується «{enc_nl}»). Print Screen — у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 3. Режим гасла",
        f"Що зробити: режим «Гасло», ключ «{motto}», текст «{plain}», "
        f"зашифруй (очікується «{enc_m}»). Print Screen — у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 4. Фрагмент коду (TrithemiusCipher)",
        "Що зробити: у crypto.py зніми Print Screen методів validate_key, encrypt, decrypt. "
        "Встав у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 5. Додаткове — атака за відомою парою",
        "Що зробити: меню «Шифрування» → «Атака за відомою парою». "
        f"Введи відкритий «{plain}» і шифротекст лінійного режиму «{enc_lin}», "
        "натисни «Знайти ключ» (має знайти 2,5). Print Screen — у рамку.",
    )

    add_center(doc, "5. Висновки", size=14)
    add_p(
        doc,
        "Реалізовано три режими ключа шифру Тритеміуса, валідацію та GUI на базі КП №1. "
        "Додатково працює модуль відновлення ключа за парою відкритий текст — шифротекст.",
    )
    add_p(doc, "Посилання на код (GitHub): ________________________________", bold=True)
    doc.save("zvit.docx")
    print("ok")


if __name__ == "__main__":
    main()
