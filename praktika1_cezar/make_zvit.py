from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from crypto import CaesarCipher, UA_ALPHABET


def set_run(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_center(doc, text, size=14, bold=True, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    return p


def add_p(doc, text, size=12, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_shot_box(doc, title, what_to_do):
    add_p(doc, title, bold=True)
    add_p(doc, what_to_do, color=RGBColor(0xC0, 0x00, 0x00))
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:color"), "C00000")
        tcBorders.append(border)
    tcPr.append(tcBorders)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "\n\n\n\n"
        ">>> ВСТАВ СЮДИ PRINT SCREEN <<<\n"
        "(Ctrl+V після PrtSc / Win+Shift+S)\n"
        "\n\n\n"
    )
    set_run(r, size=11, bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    doc.add_paragraph()


def main():
    c = CaesarCipher()
    key = 13
    plain = "Алфавіт"
    enc = c.encrypt(plain, key, UA_ALPHABET)
    dec = c.decrypt(enc, key, UA_ALPHABET)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    add_center(
        doc,
        "НАЦІОНАЛЬНИЙ ТЕХНІЧНИЙ УНІВЕРСИТЕТ УКРАЇНИ",
        size=12,
        bold=True,
    )
    add_center(doc, "«КИЇВСЬКИЙ ПОЛІТЕХНІЧНИЙ ІНСТИТУТ імені ІГОРЯ СІКОРСЬКОГО»", size=12)
    add_center(
        doc,
        "Дисципліна: Безпека інформаційних систем",
        size=12,
        color=RGBColor(0x00, 0x55, 0x99),
    )
    doc.add_paragraph()
    add_center(doc, "ЗВІТ", size=16, bold=True)
    add_center(doc, "до комп'ютерного практикуму №1", size=14, bold=True)
    add_center(doc, "Тема: Шифр Цезаря", size=14, bold=True)
    doc.add_paragraph()
    add_p(doc, "Студент: Михайлець Артем Миколайович")
    add_p(doc, "Група: ТВ-33")
    add_p(doc, "Номер у списку групи (ключ): 13")
    add_p(doc, "Викладач: ________________________________")
    add_p(doc, "м. Київ, 2026 р.")
    doc.add_page_break()

    add_center(doc, "1. Мета роботи", size=14)
    add_p(
        doc,
        "Розробити криптосистему симетричного шифрування на основі шифру Цезаря "
        "з графічним інтерфейсом, валідацією ключа та даних, шифруванням/розшифруванням "
        "українською та англійською мовами.",
    )

    add_center(doc, "2. Короткі теоретичні відомості", size=14)
    add_p(
        doc,
        "Шифр Цезаря — підстановка, у якій кожен символ зсувається на фіксовану "
        "кількість позицій k в алфавіті потужності n:",
    )
    add_p(doc, "y = (x + k) mod n", bold=True)
    add_p(doc, "x = (y + n − (k mod n)) mod n", bold=True)
    add_p(
        doc,
        "У програмі використано український алфавіт (33 літери, з ґ) та англійський (26 літер). "
        "Символи поза алфавітом (пробіли, розділові знаки) не змінюються.",
    )

    add_center(doc, "3. Хід виконання", size=14)
    add_p(doc, "3.1. Інтерфейс програми", bold=True)
    add_p(
        doc,
        "Реалізовано меню та панель інструментів: створення/відкриття/збереження/друк файлів, "
        "шифрування і розшифрування, відомості про розробника, вихід. "
        "Передбачено атаку грубою силою та шифрування бінарних файлів (по байтах, mod 256).",
    )

    add_p(doc, "3.2. Тестовий приклад (за номером у списку)", bold=True)
    add_p(doc, f'Відкритий текст: "{plain}"')
    add_p(doc, f"Ключ k = {key}")
    add_p(doc, f'Шифротекст: "{enc}"')
    add_p(doc, f'Після розшифрування: "{dec}"')

    add_center(doc, "4. Результати тестування (скріншоти)", size=14)

    add_shot_box(
        doc,
        "Крок 1. Шифрування",
        "Що зробити: запусти python main.py → введи «Алфавіт» → ключ 13 → "
        "натисни «Зашифрувати». Зроби Print Screen усього вікна програми "
        f'(має бути результат «{enc}») і встав у рамку нижче.',
    )

    add_shot_box(
        doc,
        "Крок 2. Фрагмент коду (клас CaesarCipher)",
        "Що зробити: відкрий crypto.py у редакторі, виділи методи validate_key, "
        "validate_data, encrypt, decrypt. Зроби Print Screen і встав у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 3. Розшифрування",
        f"Що зробити: вставь у поле тексту «{enc}», ключ 13, натисни «Розшифрувати». "
        f'Має вийти «{dec}». Print Screen усього вікна — у рамку нижче.',
    )

    add_shot_box(
        doc,
        "Крок 4. Додаткове завдання — атака грубою силою",
        "Що зробити: меню «Шифрування» → «Атака грубою силою». "
        "Print Screen вікна з перебором ключів — у рамку нижче.",
    )

    add_shot_box(
        doc,
        "Крок 5. Відомості про розробника",
        "Що зробити: меню «Довідка» → «Про розробника». Print Screen діалогу — у рамку.",
    )

    add_center(doc, "5. Висновки", size=14)
    add_p(
        doc,
        "У ході роботи реалізовано криптосистему на основі шифру Цезаря з GUI (tkinter), "
        "валідацією ключа/даних, підтримкою UA/EN алфавітів, роботою з файлами, "
        "атакою перебором і шифруванням довільних байтових файлів. "
        "Тестове слово «Алфавіт» з ключем 13 коректно шифрується і розшифровується.",
    )

    add_p(doc, "Посилання на код (GitHub): ________________________________", bold=True)
    doc.save("zvit.docx")
    print("ok")


if __name__ == "__main__":
    main()
