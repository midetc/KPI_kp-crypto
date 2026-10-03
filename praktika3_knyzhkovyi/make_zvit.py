import random

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from crypto import DEFAULT_POEM, VerseCipher, WordBookCipher


def set_run(run, size=12, bold=False, color=None):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_center(doc, text, size=14, bold=True, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)


def add_p(doc, text, size=12, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(6)


def add_shot_box(doc, title, what_to_do):
    add_p(doc, title, bold=True)
    add_p(doc, what_to_do, color=RGBColor(0xC0, 0x00, 0x00))
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.width = Cm(16)
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "12")
        border.set(qn("w:color"), "C00000")
        tcBorders.append(border)
    tcPr.append(tcBorders)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(
        "\n\n\n\n"
        ">>> ВСТАВ СЮДИ PRINT SCREEN <<<\n"
        "(Ctrl+V після PrtSc / Win+Shift+S)\n"
        "\n\n\n"
    )
    set_run(r, size=11, bold=True, color=RGBColor(0xC0, 0x00, 0x00))
    doc.add_paragraph()


def main():
    cipher = VerseCipher()
    word_c = WordBookCipher()
    plain = "Алфавіт"
    width = 10
    enc = cipher.encrypt(plain, DEFAULT_POEM, width, random.Random(13))
    dec = cipher.decrypt(enc, DEFAULT_POEM, width)
    grid = cipher.validate_key(DEFAULT_POEM, width).as_text()
    book = "Алла любить фіалки а весна і тиша"
    enc_w = word_c.encrypt("алів", book, random.Random(13))
    dec_w = word_c.decrypt(enc_w, book)

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(1.5)

    add_center(doc, "НАЦІОНАЛЬНИЙ ТЕХНІЧНИЙ УНІВЕРСИТЕТ УКРАЇНИ", size=12, bold=True)
    add_center(doc, "«КИЇВСЬКИЙ ПОЛІТЕХНІЧНИЙ ІНСТИТУТ імені ІГОРЯ СІКОРСЬКОГО»", size=12)
    add_center(
        doc,
        "Дисципліна: Безпека інформаційних систем",
        size=12,
        color=RGBColor(0x00, 0x55, 0x99),
    )
    doc.add_paragraph()
    add_center(doc, "ЗВІТ", size=16, bold=True)
    add_center(doc, "до комп'ютерного практикуму №3", size=14, bold=True)
    add_center(doc, "Тема: Книжковий шифр (віршований)", size=14, bold=True)
    doc.add_paragraph()
    add_p(doc, "Студент: Михайлець Артем Миколайович")
    add_p(doc, "Група: ТВ-33")
    add_p(doc, "Викладач: ________________________________")
    add_p(doc, "м. Київ, 2026 р.")
    doc.add_page_break()

    add_center(doc, "1. Мета роботи", size=14)
    add_p(
        doc,
        "Розробити криптосистему на основі використання віршованого фрагменту "
        "як ключа шифрування (книжковий / віршований шифр).",
    )

    add_center(doc, "2. Короткі теоретичні відомості", size=14)
    add_p(
        doc,
        "У віршованому шифрі ключ — вірш, записаний у прямокутну таблицю. "
        "Рядки і стовпці нумеруються. Кожній літері відкритого тексту "
        "ставять у відповідність координати такої ж літери в таблиці у форматі рядок/стовпець "
        "(наприклад 4/3). Один і той самий символ може мати кілька різних кодів.",
    )
    add_p(
        doc,
        "Додатково реалізовано різновид книжкового шифру: номер слова в тексті-книзі, "
        "перша літера якого збігається з літерою повідомлення.",
    )

    add_center(doc, "3. Ключова таблиця та тест", size=14)
    add_p(doc, f"Ширина таблиці: {width}")
    add_p(doc, "Фрагмент ключа (вірш):")
    p = doc.add_paragraph()
    r = p.add_run(grid)
    set_run(r, size=9)
    r.font.name = "Consolas"
    add_p(doc, f'Відкритий текст: "{plain}"')
    add_p(doc, f'Шифрограма: "{enc}"')
    add_p(doc, f'Після розшифрування: "{dec}"')
    add_p(doc, f'Додатково (номер слова), «алів» → "{enc_w}" → "{dec_w}"')

    add_center(doc, "4. Результати тестування (скріншоти)", size=14)

    add_shot_box(
        doc,
        "Крок 1. Ключова таблиця",
        "Що зробити: запусти python main.py → «Показати таблицю ключа». "
        "Print Screen таблиці — у рамку нижче.",
    )

    add_shot_box(
        doc,
        "Крок 2. Шифрування віршованим шифром",
        f"Що зробити: режим «Віршований», ширина 10, текст «{plain}», "
        "натисни «Зашифрувати». Print Screen усього вікна — у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 3. Розшифрування",
        "Що зробити: скопіюй шифрограму з результату в поле тексту, "
        "натисни «Розшифрувати». Має повернутися вихідне слово. Print Screen — у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 4. Фрагмент коду",
        "Що зробити: Print Screen класів VerseCipher та WordBookCipher з crypto.py — у рамку.",
    )

    add_shot_box(
        doc,
        "Крок 5. Додаткове — книжковий шифр (номер слова)",
        "Що зробити: обери режим «Книжковий (номер слова)», введи короткий текст-книгу "
        "у поле ключа, зашифруй кілька літер. Print Screen — у рамку.",
    )

    add_center(doc, "5. Висновки", size=14)
    add_p(
        doc,
        "Реалізовано віршований шифр з таблицею-ключем і форматом коди рядок/стовпець, "
        "а також додатковий книжковий варіант за номером слова. "
        "Інтерфейс підтримує файли, друк і відомості про розробника.",
    )
    add_p(doc, "Посилання на код (GitHub): ________________________________", bold=True)
    doc.save("zvit.docx")
    print("ok")


if __name__ == "__main__":
    main()
