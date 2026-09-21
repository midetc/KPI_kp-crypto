from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from crypto import CaesarCipher, TrithemiusCipher, UA_ALPHABET


def add_title(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_red(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)


def box(doc, lines):
    for line in lines:
        doc.add_paragraph(line)


def report_lab1():
    c = CaesarCipher()
    key = 13
    plain = "Алфавіт"
    enc = c.encrypt(plain, key, UA_ALPHABET)
    dec = c.decrypt(enc, key, UA_ALPHABET)

    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Безпека інформаційних систем")
    r.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
    r.bold = True

    doc.add_paragraph("Група ТВ-33")
    doc.add_paragraph("Михайлець Артем Миколайович")
    add_title(doc, "Комп'ютерний практикум №1")
    add_title(doc, "Тема: Шифр Цезаря")
    doc.add_paragraph("Мета: розробити криптосистему на основі шифру Цезаря.")

    doc.add_paragraph()
    add_title(doc, "Крок 1. Шифрування")
    box(
        doc,
        [
            'Інтерфейс програми "Шифр Цезаря" (Михайлець А.М.)',
            f'Ввід: "{plain}"',
            f"Ключ (номер у списку групи): {key}",
            f'Результат: "{enc}"',
        ],
    )
    add_red(doc, "[Сюди вставити Print Screen вікна з шифруванням]")

    doc.add_paragraph()
    add_title(doc, "Крок 2. Фрагмент коду")
    add_red(
        doc,
        "[Сюди вставити Print Screen частини лістингу: validate_key / encrypt / decrypt]",
    )

    doc.add_paragraph()
    add_title(doc, "Крок 3. Розшифрування")
    box(
        doc,
        [
            'Інтерфейс програми "Шифр Цезаря"',
            f'Ввід: "{enc}"',
            f"Ключ: {key}",
            f'Результат: "{dec}"',
        ],
    )
    add_red(doc, "[Сюди вставити Print Screen вікна з розшифруванням]")

    doc.add_paragraph()
    doc.add_paragraph(
        "Додатково: реалізовано атаку грубою силою (перебір ключів) "
        "та шифрування довільних (бінарних) файлів по байтах mod 256."
    )
    add_red(doc, "Решта лабораторних — аналогічно.")
    doc.save("zvit_lab1_cezar.docx")


def report_lab2():
    t = TrithemiusCipher()
    plain = "Алфавіт"
    linear_key = "2,5"
    enc_lin = t.encrypt(plain, "linear", linear_key, UA_ALPHABET)
    dec_lin = t.decrypt(enc_lin, "linear", linear_key, UA_ALPHABET)

    nonlin_key = "1,2,3"
    enc_nl = t.encrypt(plain, "nonlinear", nonlin_key, UA_ALPHABET)

    motto = "ключ"
    enc_m = t.encrypt(plain, "motto", motto, UA_ALPHABET)

    doc = Document()
    p = doc.add_paragraph()
    r = p.add_run("Безпека інформаційних систем")
    r.font.color.rgb = RGBColor(0x00, 0x55, 0x99)
    r.bold = True

    doc.add_paragraph("Група ТВ-33")
    doc.add_paragraph("Михайлець Артем Миколайович")
    add_title(doc, "Комп'ютерний практикум №2")
    add_title(doc, "Тема: Шифр Тритеміуса")
    doc.add_paragraph(
        "Мета: розробити криптосистему на основі шифру Тритеміуса "
        "(модифікація системи з КП №1)."
    )

    doc.add_paragraph()
    add_title(doc, "Крок 1. Лінійний ключ (A, B)")
    box(
        doc,
        [
            f'Ввід: "{plain}"',
            f"Ключ: {linear_key}  (k = A*p + B)",
            f'Шифротекст: "{enc_lin}"',
            f'Після розшифрування: "{dec_lin}"',
        ],
    )
    add_red(doc, "[Print Screen інтерфейсу для лінійного режиму]")

    doc.add_paragraph()
    add_title(doc, "Крок 2. Нелінійний ключ (A, B, C)")
    box(
        doc,
        [
            f'Ввід: "{plain}"',
            f"Ключ: {nonlin_key}  (k = A*p^2 + B*p + C)",
            f'Шифротекст: "{enc_nl}"',
        ],
    )
    add_red(doc, "[Print Screen інтерфейсу для нелінійного режиму]")

    doc.add_paragraph()
    add_title(doc, "Крок 3. Гасло")
    box(
        doc,
        [
            f'Ввід: "{plain}"',
            f'Гасло: "{motto}"',
            f'Шифротекст: "{enc_m}"',
        ],
    )
    add_red(doc, "[Print Screen інтерфейсу для режиму з гаслом]")

    doc.add_paragraph()
    add_title(doc, "Крок 4. Код")
    add_red(
        doc,
        "[Print Screen класів TrithemiusCipher та KnownPlaintextAttack]",
    )

    doc.add_paragraph()
    doc.add_paragraph(
        "Додаткове завдання: модуль активної атаки відновлює ключ "
        "за парою відкритий текст — шифротекст (меню Шифрування → "
        "Атака за відомою парою)."
    )
    doc.save("zvit_lab2_trytemius.docx")


if __name__ == "__main__":
    report_lab1()
    report_lab2()
    print("ok")
